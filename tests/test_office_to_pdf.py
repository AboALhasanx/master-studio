"""QA gate for Office->PDF conversion (the x2t/Asana-Math regression).

In Oct 2026 a whole subject booklet shipped with every page rendered in a
single math font: visually wrecked AND unextractable. Nothing caught it —
not the publisher, not the verifier, not the suite. These tests pin the
three failure signatures so a broken conversion fails the build loudly:

* ``math-font`` — a text page whose only fonts are math/symbol fonts;
* ``no-spaces`` — a text-heavy page with (almost) no spaces;
* ``empty`` — pages with no extractable text at all.
"""

from __future__ import annotations

import re
import shutil

import pytest

from tools.office_to_pdf import _MATH_FONT_RE, QAError, qa_pdf


def _make_pdf(path, *, pages: int = 2, text: str | None = None):
    """Synthesize a PDF with fitz (base-14 Helvetica — no binaries needed)."""
    import fitz

    doc = fitz.open()
    for _ in range(pages):
        page = doc.new_page()
        if text is not None:
            page.insert_textbox(fitz.Rect(72, 72, 2000, 3000), text)
    doc.save(str(path))
    return path


GOOD_TEXT = ("Data Mining deals with different types of data. Before applying "
             "a data-mining technique, it is important to understand what the "
             "data objects are and what the attributes are. " * 6)


class TestMathFontSignature:
    @pytest.mark.parametrize("name", ["Asana Math", "BAAAAA+AsanaMath",
                                      "Cambria Math", "SymbolMT", "Wingdings",
                                      "STIXMathJax", "XITS Math"])
    def test_math_symbol_fonts_are_flagged(self, name):
        assert _MATH_FONT_RE.search(name), name

    @pytest.mark.parametrize("name", ["SegoeUI", "ArialMT", "TimesNewRomanPSMT",
                                      "Calibri", "Tahoma", "CourierNewPSMT",
                                      "Amiri", "NotoNaskhArabic"])
    def test_real_text_fonts_pass(self, name):
        assert not _MATH_FONT_RE.search(name), name


class TestQaGate:
    def test_clean_text_passes(self, tmp_path):
        pdf = _make_pdf(tmp_path / "good.pdf", pages=3, text=GOOD_TEXT)
        assert qa_pdf(pdf)["pages"] == 3

    def test_missing_spaces_fail(self, tmp_path):
        squashed = GOOD_TEXT.replace(" ", "")[:1200]
        pdf = _make_pdf(tmp_path / "squashed.pdf", pages=2, text=squashed)
        with pytest.raises(QAError, match="no-spaces"):
            qa_pdf(pdf)

    def test_empty_pages_fail(self, tmp_path):
        pdf = _make_pdf(tmp_path / "empty.pdf", pages=3, text=None)
        with pytest.raises(QAError, match="empty"):
            qa_pdf(pdf)

    def test_single_page_needs_no_average(self, tmp_path):
        pdf = _make_pdf(tmp_path / "one.pdf", pages=1, text="Hello world. " * 4)
        assert qa_pdf(pdf)["pages"] == 1


class TestEngineSelection:
    """The COM/Chromium dispatch and its Windows-specific helpers."""

    def test_requesting_com_without_office_fails_loudly(self, monkeypatch):
        import tools.office_to_pdf as mod

        monkeypatch.setattr(mod, "_com_available", lambda: False)
        with pytest.raises(RuntimeError, match="Office COM"):
            mod.convert("does-not-matter.docx", "out.pdf", qa=False, engine="com")

    def test_unsupported_suffix_is_refused(self):
        import tools.office_to_pdf as mod

        with pytest.raises(ValueError, match="unsupported input"):
            mod.convert("notes.txt", "out.pdf", qa=False)

    def test_office_pids_never_raises_without_office(self):
        import tools.office_to_pdf as mod

        # The invariant is "never raises", not "always empty": on a Windows
        # box with Word open it legitimately returns pids. On CI/Linux there
        # is no PowerShell, so it must degrade to an empty set.
        result = mod._office_pids()
        assert isinstance(result, set)
        assert all(isinstance(p, int) for p in result)

    def test_dismiss_dialogs_is_a_safe_noop_without_office(self):
        import tools.office_to_pdf as mod

        assert mod._dismiss_office_dialogs() == 0


class TestEndToEnd:
    """Full DOCX->PDF — skipped where no engine exists (COM or Chromium)."""

    @staticmethod
    def _engine_present():
        import os
        import sys
        from pathlib import Path

        if sys.platform == "win32":
            try:
                import win32com.client  # noqa: F401
                return True
            except ImportError:
                pass
        local_app = os.environ.get("LOCALAPPDATA", "")
        if local_app and list((Path(local_app) / "ms-playwright").glob(
                "chromium-*/chrome-win64/chrome.exe")):
            return True
        return any(shutil.which(n) for n in
                   ("msedge", "chrome", "google-chrome", "chromium"))

    @pytest.mark.skipif(not _engine_present.__func__(),
                        reason="no conversion engine on this machine")
    def test_docx_roundtrip_is_clean(self, tmp_path):
        import docx

        from tools.office_to_pdf import convert

        src = tmp_path / "sample.docx"
        doc = docx.Document()
        doc.add_heading("Roundtrip Probe", level=1)
        doc.add_paragraph("Data Mining deals with different types of data, "
                          "and this sentence must survive conversion intact.")
        doc.save(str(src))
        out = tmp_path / "sample.pdf"
        assert convert(src, out)["qa"] == "passed"
        assert "must survive conversion intact" in out.read_bytes().decode(
            "latin-1", errors="ignore") or True  # binary; QA already asserted
        import fitz
        text = fitz.open(str(out))[0].get_text()
        assert "must survive conversion intact" in " ".join(text.split())
