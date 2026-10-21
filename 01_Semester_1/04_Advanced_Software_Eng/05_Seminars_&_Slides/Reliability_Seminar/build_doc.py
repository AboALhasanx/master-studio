#!/usr/bin/env python3
"""Build the Chapter 7 (Software Reliability) Word document - Part 1."""
import re, sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = Path(__file__).resolve().parent
FIG = HERE / "figures"
MD = HERE / "ch7_part1.md"
OUT = HERE / "Ch7_Software_Reliability.docx"

FONT = "Times New Roman"
NAVY = RGBColor(0x00, 0x00, 0x00)
GREY = RGBColor(0x44, 0x44, 0x44)


def set_base(doc):
    n = doc.styles["Normal"]
    n.font.name = FONT
    n.font.size = Pt(11)
    n.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
    n.paragraph_format.space_after = Pt(6)
    n.paragraph_format.line_spacing = 1.15
    rpr = n.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), FONT)


def add_field(p, instr):
    r = p.add_run()
    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin")
    i = OxmlElement("w:instrText"); i.set(qn("xml:space"), "preserve"); i.text = instr
    s = OxmlElement("w:fldChar"); s.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t"); t.text = "Update field (F9)"
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end")
    for el in (b, i, s, t, e): r._r.append(el)


def runs(p, text, size=11, italic=False, bold=False):
    toks = re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*)', text)
    for tok in toks:
        if not tok: continue
        if tok.startswith("**") and tok.endswith("**") and len(tok) >= 4:
            r = p.add_run(tok[2:-2]); r.bold = True
        elif tok.startswith("*") and tok.endswith("*") and len(tok) >= 2:
            r = p.add_run(tok[1:-1]); r.italic = True
        else:
            r = p.add_run(tok)
        r.font.name = FONT; r.font.size = Pt(size); r.italic = italic or r.italic
    return p


def heading(doc, text, level):
    h = doc.add_heading(level=level)
    r = h.add_run(text)
    r.font.name = FONT; r.bold = True
    r.font.color.rgb = NAVY
    r.font.size = Pt(16 if level == 1 else 13 if level == 2 else 11.5)
    h.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    h.paragraph_format.space_after = Pt(6)
    return h


def add_image(doc, fname, width, caption):
    from docx.enum.text import WD_BREAK
    path = FIG / fname
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.keep_together = True
    if path.exists():
        r = p.add_run(); r.add_picture(str(path), width=Inches(width))
    else:
        runs(p, f"[missing figure: {fname}]")
    if caption:
        br = p.add_run(); br.add_break(WD_BREAK.LINE)
        c = p.add_run(caption)
        c.italic = True; c.font.size = Pt(9); c.font.color.rgb = GREY; c.font.name = FONT


def add_table(doc, rows):
    if not rows: return
    ncol = max(len(r) for r in rows)
    t = doc.add_table(rows=0, cols=ncol); t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    lens = [max((len(str(r[c])) if c < len(r) else 0) for r in rows) for c in range(ncol)]
    total = sum(lens) or 1
    widths = [max(0.55, min(3.4, (m / total) * 6.3)) for m in lens]
    scale = 6.3 / sum(widths)
    widths = [w * scale for w in widths]
    for i, row in enumerate(rows):
        cells = t.add_row().cells
        # let rows break across pages normally; no header repetition
        trPr = t.rows[-1]._tr.get_or_add_trPr()
        cs = OxmlElement("w:cantSplit"); cs.set(qn("w:val"), "0"); trPr.append(cs)
        for c in range(ncol):
            cells[c].width = Inches(widths[c])
            val = row[c] if c < len(row) else ""
            cell = cells[c]; cell.text = ""
            par = cell.paragraphs[0]; par.paragraph_format.space_after = Pt(1)
            runs(par, val, size=9)
            if i == 0:
                for run in par.runs:
                    run.bold = True; run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), "3A3A3A")
                cell._element.get_or_add_tcPr().append(shd)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def cover(doc):
    for _ in range(2): doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Software Reliability"); r.font.name = FONT; r.font.size = Pt(30); r.bold = True; r.font.color.rgb = NAVY
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Chapter 7 — Advanced Software Engineering"); r.font.name = FONT; r.font.size = Pt(14); r.font.color.rgb = GREY
    doc.add_paragraph()
    meta = [
        ("University", "University of Wasit"),
        ("College", "College of Computer Science and Information Technology"),
        ("Prepared by", "Abul-Hassan Abbas Khudair"),
        ("Supervised by", "Asst. Prof. Dr. Ali Fahim Ni'ma"),
        ("Year", "2026–2027"),
    ]
    for k, v in meta:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        r1 = p.add_run(f"{k}:  "); r1.bold = True; r1.font.name = FONT; r1.font.size = Pt(11)
        r2 = p.add_run(v); r2.font.name = FONT; r2.font.size = Pt(11)
    doc.add_page_break()
    # TOC title as a plain styled paragraph (so it is not listed inside the TOC)
    cp = doc.add_paragraph()
    r = cp.add_run("Table of Contents")
    r.font.name = FONT; r.bold = True; r.font.size = Pt(16); r.font.color.rgb = NAVY
    cp.paragraph_format.space_before = Pt(14); cp.paragraph_format.space_after = Pt(6)
    p = doc.add_paragraph(); add_field(p, r'TOC \o "1-3" \h \z \u')
    doc.add_page_break()


def parse(doc):
    lines = MD.read_text(encoding="utf-8").split("\n")
    i = 0
    while i < len(lines):
        s = lines[i].strip(); i += 1
        m = re.match(r'^!\[(.+?)\|([\d.]+)\]\s*$', s)
        if m:
            cap = ""
            if i < len(lines) and lines[i].strip().startswith("*Figure"):
                cap = lines[i].strip().strip("*"); i += 1
            add_image(doc, m.group(1), float(m.group(2)), cap)
            continue
        if not s or s == "---": continue
        if s.startswith("# "): heading(doc, s[2:], 1)
        elif s.startswith("## "): heading(doc, s[3:], 2)
        elif s.startswith("### "): heading(doc, s[4:], 3)
        elif s.startswith("> "):
            p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.35)
            runs(p, s[2:], size=10, italic=True); [setattr(r.font, "color", GREY) for r in p.runs]
        elif s.startswith("- "):
            p = doc.add_paragraph(style="List Bullet"); p.paragraph_format.space_after = Pt(2)
            runs(p, s[2:])
        elif re.match(r'^\d+\.\s', s):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.35)
            p.paragraph_format.first_line_indent = Inches(-0.22)
            p.paragraph_format.space_after = Pt(2)
            runs(p, s)
        elif s.startswith("|") and s.endswith("|"):
            block = [s]
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i].strip()); i += 1
            rows = []
            for br in block:
                if re.match(r'^\|[\s\-:|]+\|$', br): continue
                rows.append([c.strip() for c in br.split("|")[1:-1]])
            add_table(doc, rows)
        else:
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_after = Pt(7)
            if re.search(r"\(7\.\d+\)\s*$", s) and "=" in s:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(6)
                p.paragraph_format.space_after = Pt(8)
                runs(p, s, size=10.5)
            else:
                runs(p, s)


def main():
    global OUT
    if len(sys.argv) > 1:
        OUT = Path(sys.argv[1])
    doc = Document(); set_base(doc)
    for s in doc.sections:
        s.top_margin = Inches(1); s.bottom_margin = Inches(1)
        s.left_margin = Inches(1); s.right_margin = Inches(1)
    cover(doc)
    parse(doc)
    # metadata: neutral
    cp = doc.core_properties
    cp.author = "-"
    cp.last_modified_by = "-"
    cp.title = "Software Reliability — Chapter 7"
    cp.subject = "-"
    cp.comments = "-"
    cp.category = "-"
    doc.save(OUT)
    print("Saved:", OUT)
    print("Paragraphs:", len(doc.paragraphs), "Tables:", len(doc.tables))


if __name__ == "__main__":
    main()
