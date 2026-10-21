#!/usr/bin/env python3
"""Five-cycle QA for the Reliability deliverable."""
import re
import sys
import zipfile
import hashlib
from pathlib import Path

import pymupdf
import numpy as np
from docx import Document
from docx.shared import Emu

BASE = Path(__file__).resolve().parent
DOCX = BASE / "Ch7_Software_Reliability.docx"
PDF = BASE / "Ch7_Software_Reliability.pdf"
MD = BASE / "ch7_part1.md"
FIG = BASE / "figures"

fails = []
warns = []


def ok(cond, label, detail=""):
    if cond:
        print(f"  PASS  {label}")
    else:
        print(f"  FAIL  {label}   {detail}")
        fails.append(label)


print("\n=== CYCLE 1: build integrity ===")
doc = Document(DOCX)
md = MD.read_text(encoding="utf-8")
words = len(md.split())
print(f"  info  paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} words={words}")
ok(len(doc.tables) == 19, "19 tables present", len(doc.tables))
ok(words > 15000, "content expanded (>15k words)", words)
text_all = "\n".join(p.text for p in doc.paragraphs)
ok("[missing figure" not in text_all, "no missing-figure placeholder")
with zipfile.ZipFile(DOCX) as z:
    media = [n for n in z.namelist() if n.startswith("word/media/")]
ok(len(media) == 24, "24 embedded images", len(media))

print("\n=== CYCLE 2: content hygiene ===")
for bad in ["**", "[THIN]", "**EN.**", "**AR.**", "TODO", "File 01 of", "BUILD_PLAN"]:
    ok(bad not in text_all, f"no leaked marker {bad!r}")
ok("```" not in md, "no ASCII-art code fences in source")
heads = [p.text.strip() for p in doc.paragraphs if p.style.name.startswith("Heading")]
print(f"  info  headings={len(heads)}")
ok(all(re.match(r"^\d(\.\d+){0,3}\s", h) for h in heads[1:]), "all headings numbered")
ok(len(heads) >= 20, "TOC-sized heading set", len(heads))

print("\n=== CYCLE 3: figures ===")
md_imgs = re.findall(r"!\[([^\]]+?)\|", md)
ok(len(md_imgs) == 24, "24 figure references in source", len(md_imgs))
missing = [f for f in md_imgs if not (FIG / f).exists()]
ok(not missing, "all referenced figures exist", missing)
hashes, coloured = {}, []
for f in md_imgs:
    pm = pymupdf.Pixmap(str(FIG / f))
    arr = np.frombuffer(pm.samples, dtype=np.uint8).reshape(pm.height, pm.width, pm.n)[:, :, :3].astype(int)
    dev = np.max(np.abs(arr - np.repeat(arr.mean(axis=2, keepdims=True), 3, axis=2)), axis=2)
    if (dev > 6).sum() > 20:
        coloured.append(f)
    h = hashlib.md5((FIG / f).read_bytes()).hexdigest()
    hashes.setdefault(h, []).append(f)
ok(not coloured, "all figures grayscale", coloured)
ok(all(len(v) == 1 for v in hashes.values()), "no duplicate figure files",
   [v for v in hashes.values() if len(v) > 1])
caps = re.findall(r"^\*?(Figure 7\.\d+\.)", md, re.M)
ok(len(caps) == 24, "24 figure captions", len(caps))
ok(len(set(caps)) == 24, "captions numbered uniquely")
nums = [int(c.split()[1].rstrip(".").split(".")[1]) for c in caps]
ok(nums == sorted(nums), "figures numbered in document order", nums)

print("\n=== CYCLE 4: layout ===")
pdf = pymupdf.open(PDF)
ok(pdf.page_count >= 40, "page count >= 40", pdf.page_count)
overflow = []
for i, page in enumerate(pdf):
    w, h = page.rect.width, page.rect.height
    for b in page.get_text("blocks"):
        x0, y0, x1, y1 = b[:4]
        if x0 < 35 or x1 > w - 35 or y0 < 25 or y1 > h - 25:
            overflow.append((i + 1, round(x0, 1), round(x1, 1)))
ok(not overflow, "no text outside margins", overflow[:4])
with zipfile.ZipFile(DOCX) as z:
    xml = z.read("word/document.xml").decode("utf-8")
ok("tblHeader" not in xml, "table header not repeated on page breaks")
blocked = len(re.findall(r'<w:cantSplit(?![^>]*w:val="(?:0|false)")', xml))
ok(blocked == 0, "no table row forced to stay unbroken", blocked)

# TOC page numbers vs real heading pages
toc_lines = [l for l in pdf[1].get_text().split("\n") if re.match(r"^7(\.\d+)*\s", l.strip())]
page_text = [p.get_text() for p in pdf]
bad_toc = []
for line in toc_lines:
    m = re.match(r"^(.*?)[\s.]{2,}(\d+)$", line.strip())
    if not m:
        continue
    title, num = m.group(1).strip(), int(m.group(2))
    hit = next((i + 3 for i, t in enumerate(page_text[2:]) if title[:40] in t.replace("\n", " ")), None)
    if hit != num:
        bad_toc.append((title[:45], num, hit))
ok(not bad_toc, "TOC page numbers match headings", bad_toc[:4])

print("\n=== CYCLE 5: fidelity & metadata ===")
cp = doc.core_properties
ok(cp.author == "-" and cp.last_modified_by == "-", "metadata neutral", (cp.author, cp.last_modified_by))
with zipfile.ZipFile(DOCX) as z:
    settings = z.read("word/settings.xml").decode("utf-8")
ok("updateFields" not in settings, "no update-fields prompt on open")

probes = [
    "Software reliability means operational reliability",
    "MTBF = MTTF + MTTR",
    "90 percent of the execution time",
    "The relationship between quality factors and quality criteria is given in Fig. 7.10",
    "Fq = c1 m1 + c2 m2",
    "ISO 9126 model was revised and replaced",
    "PS = PA + PR",
    "Software may be retired only if it becomes obsolete",
]
flat = re.sub(r"\s+", " ", text_all)
for p in probes:
    ok(p in flat, f"source text present: {p[:48]}")

print("\n================ SUMMARY ================")
print(f"FAILURES: {len(fails)}")
for f in fails:
    print("  -", f)
sys.exit(1 if fails else 0)
