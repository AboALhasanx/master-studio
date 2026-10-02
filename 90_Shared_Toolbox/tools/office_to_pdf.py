"""Office -> PDF converter: pip-only, no SaaS, no LibreOffice, no x2t.

Pipeline
--------
DOCX --(mammoth)--> HTML --(Chromium print-to-PDF)--> PDF
PPTX --(python-pptx)--> HTML slides --(Chromium)--> PDF (landscape)

Why this exists: the ONLYOFFICE ``x2t`` converter rendered whole documents in
a single math font (Asana Math) — every published chapter built with it is
visually wrecked *and* extracts garbage. This module replaces x2t everywhere
(see ``TELEGRAM_PUBLISHING_PLAYBOOK.md``).

The :func:`qa_pdf` gate encodes that failure signature so no broken PDF can
ever ship again: a build whose pages use only math/symbol fonts, whose text
has (almost) no spaces, or whose pages are empty FAILS loudly (exit 2).

Usage::

    python 90_Shared_Toolbox/tools/office_to_pdf.py in.docx -o out.pdf
    python 90_Shared_Toolbox/tools/office_to_pdf.py slides.pptx -o out.pdf
    python 90_Shared_Toolbox/tools/office_to_pdf.py in.docx -o out.pdf --no-qa
"""

from __future__ import annotations

import argparse
import base64
import io
import re
import sys
from pathlib import Path

__all__ = ["convert", "qa_pdf", "QAError"]

#: Font names that must never be a page's *only* font when the page carries
#: text — the exact x2t failure signature (whole document in Asana Math).
_MATH_FONT_RE = re.compile(r"math|symbol|wingding|webding|asana|cambria",
                           re.IGNORECASE)


class QAError(Exception):
    """Raised when a converted PDF fails the text sanity gate."""


# --------------------------------------------------------------------------
# DOCX -> HTML
# --------------------------------------------------------------------------
def _docx_to_html(source: Path, workdir: Path) -> str:
    import mammoth

    images = {}

    def _save_image(image):
        name = f"img{len(images)}{image.content_type.rpartition('/')[2] or '.png'}"
        data = image.read()
        (workdir / name).write_bytes(data)
        images[name] = True
        return {"src": name}

    with open(source, "rb") as fh:
        result = mammoth.convert_to_html(fh, convert_image=mammoth.images.img_element(_save_image))
    return result.value


# --------------------------------------------------------------------------
# PPTX -> HTML
# --------------------------------------------------------------------------
def _emu_to_px(emu: int) -> float:
    return emu / 914400 * 96.0


def _pptx_to_html(source: Path, workdir: Path) -> tuple[str, bool]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = Presentation(str(source))
    slide_w = _emu_to_px(prs.slide_width)
    slide_h = _emu_to_px(prs.slide_height)
    parts = []
    for idx, slide in enumerate(prs.slides):
        shapes = []
        for shape in slide.shapes:
            shapes.append(_shape_html(shape, workdir, 0, 0))
        parts.append(
            f'<section class="slide" style="width:{slide_w:.0f}px;'
            f'min-height:{slide_h:.0f}px">{"".join(shapes)}</section>'
        )
    return "\n".join(parts), True


def _run_text(run) -> str:
    import html as _html
    text = _html.escape(run.text or "")
    if not text:
        return ""
    css = []
    if run.font.bold:
        css.append("font-weight:bold")
    if run.font.italic:
        css.append("font-style:italic")
    size = getattr(run.font.size, "pt", None) if run.font.size else None
    if size:
        css.append(f"font-size:{size:.1f}pt")
    color = getattr(getattr(run.font, "color", None), "rgb", None)
    if color is not None:
        css.append(f"color:#{color}")
    style = f' style="{";".join(css)}"' if css else ""
    return f"<span{style}>{text}</span>"


def _shape_html(shape, workdir: Path, dx: float, dy: float) -> str:
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        return "".join(_shape_html(s, workdir, dx + _emu_to_px(shape.left),
                                   dy + _emu_to_px(shape.top))
                       for s in shape.shapes)
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        ext = (shape.image.ext or "png").split("/")[-1]
        name = f"slideimg{abs(id(shape)) % 10_000_000}.{ext}"
        (workdir / name).write_bytes(shape.image.blob)
        w = _emu_to_px(shape.width)
        return f'<img src="{name}" style="width:{w:.0f}px"/>'
    if shape.has_table:
        rows = []
        for row in shape.table.rows:
            cells = "".join(f"<td>{''.join(_run_text(r) for p in cell.text_frame.paragraphs for r in p.runs)}</td>"
                            for cell in row.cells)
            rows.append(f"<tr>{cells}</tr>")
        return f'<table class="ptable">{"".join(rows)}</table>'
    if shape.has_text_frame:
        paras = []
        for para in shape.text_frame.paragraphs:
            parts = str(para.alignment or "").split()
            align = parts[-1].lower() if parts else ""
            css = ""
            if "center" in align:
                css = ' style="text-align:center"'
            elif "right" in align:
                css = ' style="text-align:right"'
            paras.append(f"<p{css}>{''.join(_run_text(r) for r in para.runs)}</p>")
        return f"<div>{''.join(paras)}</div>"
    return ""


# --------------------------------------------------------------------------
# HTML -> PDF (Chromium)
# --------------------------------------------------------------------------
_PAGE_CSS = """
@page { size: A4; margin: 18mm 15mm; }
@page landscape-slide { size: A4 landscape; margin: 10mm; }
body { font-family: "Segoe UI", Tahoma, Arial, sans-serif; line-height: 1.6;
       color: #111; direction: ltr; }
body.rtl { direction: rtl; }
h1, h2, h3 { color: #0b3d62; }
table { border-collapse: collapse; margin: 0.6em 0; }
td, th { border: 1px solid #999; padding: 4px 8px; }
img { max-width: 100%; }
.slide { page: landscape-slide; page-break-after: always; overflow: hidden; }
section.slide:last-child { page-break-after: avoid; }
.ptable td { font-size: 11pt; }
"""

_HTML_SHELL = """<!DOCTYPE html><html><head><meta charset="utf-8"/>
<style>{css}</style></head><body>{body}</body></html>"""


def _find_chromium_binary() -> str | None:
    """Locate a launchable Chromium, mirroring diagram_forge.py's resolver.

    Playwright's *headless shell* build cannot be launched via
    ``chromium.launch()`` in every version (and the pinned shell may be
    missing while a full ``chromium-*/chrome-win64/chrome.exe`` sits on
    disk), so prefer the full build, then Edge/Chrome.
    """
    import os
    import shutil

    local_app = os.environ.get("LOCALAPPDATA", "")
    if local_app:
        pw_dir = Path(local_app) / "ms-playwright"
        if pw_dir.exists():
            for chrome_exe in sorted(
                    pw_dir.glob("chromium-*/chrome-win64/chrome.exe"),
                    reverse=True):
                if chrome_exe.exists():
                    return str(chrome_exe)
    pf = os.environ.get("ProgramFiles", r"C:\Program Files")
    pf86 = os.environ.get("ProgramFiles(x86)", r"C:\Program Files (x86)")
    for c in [Path(pf86) / "Microsoft" / "Edge" / "Application" / "msedge.exe",
              Path(pf) / "Microsoft" / "Edge" / "Application" / "msedge.exe",
              Path(pf) / "Google" / "Chrome" / "Application" / "chrome.exe",
              Path(pf86) / "Google" / "Chrome" / "Application" / "chrome.exe"]:
        if c.exists():
            return str(c)
    for name in ["msedge", "chrome", "google-chrome", "chromium"]:
        p = shutil.which(name)
        if p and Path(p).exists():
            return p
    return None


def _html_to_pdf(html: str, workdir: Path, out: Path) -> None:
    from playwright.sync_api import sync_playwright

    (workdir / "doc.html").write_text(html, encoding="utf-8")
    chromium = _find_chromium_binary()
    launch_kwargs: dict = {}
    if chromium:
        launch_kwargs["executable_path"] = chromium
    with sync_playwright() as pw:
        browser = pw.chromium.launch(**launch_kwargs)
        try:
            page = browser.new_page()
            page.goto((workdir / "doc.html").as_uri(),
                      wait_until="networkidle")
            page.pdf(path=str(out), format="A4",
                     print_background=True, prefer_css_page_size=True)
        finally:
            browser.close()


# --------------------------------------------------------------------------
# QA gate
# --------------------------------------------------------------------------
def qa_pdf(path: Path) -> dict:
    """Inspect a converted PDF; raise :class:`QAError` on failure.

    Three checks, each a standalone reason (the message says which)::

    * ``math-font`` — a page carrying real text whose *only* fonts match
      :data:`_MATH_FONT_RE` (the x2t/Asana-Math signature);
    * ``no-spaces`` — a text-heavy page with (almost) no spaces;
    * ``empty`` — a multi-page document averaging < 50 chars/page.
    """
    import fitz

    doc = fitz.open(str(path))
    problems = []
    total_chars = 0
    for i, page in enumerate(doc):
        text = page.get_text()
        total_chars += len(text.strip())
        if len(text.strip()) < 50:
            continue
        fonts = {f[3] for f in page.get_fonts()}
        if fonts and all(_MATH_FONT_RE.search(name or "") for name in fonts):
            problems.append(f"page {i + 1}: math-font "
                            f"(only fonts: {sorted(fonts)})")
        if len(text) > 200 and text.count(" ") / len(text) < 0.03:
            problems.append(f"page {i + 1}: no-spaces "
                            f"(spaces={text.count(' ')}, chars={len(text)})")
    if doc.page_count > 1 and total_chars / doc.page_count < 50:
        problems.append(f"empty (avg {total_chars / doc.page_count:.0f} chars/page)")
    if not total_chars:
        problems.append("empty (no extractable text at all)")
    if problems:
        raise QAError("; ".join(problems))
    return {"pages": doc.page_count, "chars": total_chars}


# --------------------------------------------------------------------------
# driver
# --------------------------------------------------------------------------
def convert(source: Path, out: Path, *, qa: bool = True,
            workdir: Path | None = None) -> dict:
    """Convert one Office file to PDF (with QA unless ``qa=False``)."""
    import shutil
    import tempfile

    source = Path(source)
    out = Path(out)
    tmp = Path(workdir) if workdir else Path(tempfile.mkdtemp(prefix="office2pdf_"))
    tmp.mkdir(parents=True, exist_ok=True)
    try:
        suffix = source.suffix.lower()
        if suffix == ".docx":
            body = _docx_to_html(source, tmp)
            html = _HTML_SHELL.format(css=_PAGE_CSS, body=body)
        elif suffix == ".pptx":
            body, _ = _pptx_to_html(source, tmp)
            html = _HTML_SHELL.format(css=_PAGE_CSS, body=body)
        else:
            raise ValueError(f"unsupported input: {suffix} (docx/pptx only)")
        _html_to_pdf(html, tmp, out)
    finally:
        if workdir is None:
            shutil.rmtree(tmp, ignore_errors=True)
    if qa:
        qa_pdf(out)
    return {"pdf": str(out), "qa": "passed" if qa else "skipped"}


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        prog="office_to_pdf.py",
        description="Convert DOCX/PPTX to PDF (pip-only; QA-gated).")
    parser.add_argument("input", help="source .docx or .pptx")
    parser.add_argument("-o", "--output", required=True, help="output .pdf")
    parser.add_argument("--no-qa", action="store_true",
                        help="skip the text sanity gate (not recommended)")
    args = parser.parse_args(argv)
    try:
        result = convert(args.input, args.output, qa=not args.no_qa)
    except QAError as exc:
        print(f"QA FAILED: {exc}", file=sys.stderr)
        return 2
    except Exception as exc:  # noqa: BLE001 - CLI must report, not trace
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1
    print(f"OK: {result['pdf']} (qa={result['qa']})")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
