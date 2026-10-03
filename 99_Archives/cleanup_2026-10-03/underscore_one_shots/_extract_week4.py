from docx import Document
import zipfile

path = r"C:\Users\gokoq\Master-Studio\01_Semester_1\03_Data_Mining\02_Raw_Materials\Week4_Feature_Selection.docx"
d = Document(path)
print("paras", len(d.paragraphs), "tables", len(d.tables))
print("=" * 60)
for i, p in enumerate(d.paragraphs):
    t = p.text.strip()
    if t:
        style = p.style.name if p.style else ""
        print(f"[{i}|{style}] {t}")
print("=" * 60, "TABLES", len(d.tables))
for ti, tbl in enumerate(d.tables):
    print(f"--- table {ti}: {len(tbl.rows)}x{len(tbl.columns)}")
    for r in tbl.rows:
        cells = [c.text.strip().replace("\n", " / ")[:120] for c in r.cells]
        print(" || ".join(cells))
print("=" * 60, "IMAGES")
with zipfile.ZipFile(path) as z:
    for n in z.namelist():
        if n.startswith("word/media/"):
            info = z.getinfo(n)
            print(n, info.file_size)
